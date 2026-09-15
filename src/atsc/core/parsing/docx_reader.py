"""DOCX extraction and structural checks (resume_parsing_spec.md §2 step 2, §3).

Body paragraphs and tables are read in document order. Text inside tables,
text boxes and headers is still extracted (so the content is scored) and the
corresponding check explains that a real ATS may drop or scramble it.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterator
from dataclasses import dataclass, field

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from atsc.core.models import CheckName, ParseabilityCheck
from atsc.core.parsing.checks import make_check

_SYMBOL_FONT_RE = re.compile(
    r"wingding|webding|symbol|dingbat|fontawesome|icomoon|material.?icons", re.IGNORECASE
)
_CONTACT_RE = re.compile(r"@|\(?\+?\d[\d\s().-]{7,}\d")
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class DocxExtraction:
    text: str
    checks: dict[CheckName, ParseabilityCheck] = field(default_factory=dict)


def _iter_body(doc: DocumentObject) -> Iterator[Paragraph | Table]:
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _textbox_texts(doc: DocumentObject) -> list[str]:
    out: list[str] = []
    for txbx in doc.element.body.iter(f"{_W}txbxContent"):
        paras = ["".join(t.text or "" for t in p.iter(f"{_W}t")) for p in txbx.iter(f"{_W}p")]
        out.extend(p for p in paras if p.strip())
    return out


def extract_docx(data: bytes) -> DocxExtraction:
    doc = Document(io.BytesIO(data))
    checks: dict[CheckName, ParseabilityCheck] = {}
    lines: list[str] = []
    last_heading = ""
    table_notes: list[str] = []

    for block in _iter_body(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            lines.append(text)
            style = (block.style.name or "").lower() if block.style is not None else ""
            if text and len(text) <= 40 and (text.isupper() or style.startswith("heading")):
                last_heading = text
        else:
            rows = block.rows
            table_notes.append(
                f"{len(rows)}x{len(block.columns)} table after '{last_heading or 'start'}'"
            )
            for row in rows:
                cells = [c.text.strip() for c in row.cells]
                if len(cells) >= 2 and cells[0] and not cells[0].endswith(":"):
                    lines.append(f"{cells[0]}: {', '.join(c for c in cells[1:] if c)}")
                else:
                    lines.extend(c for c in cells if c)

    boxes = _textbox_texts(doc)
    body_text = "\n".join(lines)

    # header / footer
    hf_hit: str | None = None
    for i, section in enumerate(doc.sections, start=1):
        for zone, part in (("header", section.header), ("footer", section.footer)):
            text = "\n".join(p.text for p in part.paragraphs).strip()
            if text and (_CONTACT_RE.search(text) or len(text) > 20):
                hf_hit = f'Content in the document {zone} (section {i}): "{text[:60]}"'
                lines.insert(0, text)
                break
        if hf_hit:
            break
    checks[CheckName.HEADER_FOOTER_CONTENT] = (
        make_check(CheckName.HEADER_FOOTER_CONTENT, "fail", hf_hit)
        if hf_hit
        else make_check(CheckName.HEADER_FOOTER_CONTENT, "pass")
    )

    checks[CheckName.TABLE_DETECTED] = (
        make_check(CheckName.TABLE_DETECTED, "fail", "; ".join(table_notes))
        if table_notes
        else make_check(CheckName.TABLE_DETECTED, "pass")
    )

    checks[CheckName.TEXT_BOX_DETECTED] = (
        make_check(
            CheckName.TEXT_BOX_DETECTED, "fail", f'Text box containing: "{" / ".join(boxes)[:80]}"'
        )
        if boxes
        else make_check(CheckName.TEXT_BOX_DETECTED, "pass")
    )

    multi = any(
        int(cols.get(qn("w:num"), "1")) >= 2
        for section in doc.sections
        for cols in section._sectPr.iter(f"{_W}cols")
    )
    checks[CheckName.MULTI_COLUMN_LAYOUT] = (
        make_check(
            CheckName.MULTI_COLUMN_LAYOUT, "fail", "Section formatted with two or more columns"
        )
        if multi
        else make_check(CheckName.MULTI_COLUMN_LAYOUT, "pass")
    )

    fonts = {run.font.name for p in doc.paragraphs for run in p.runs if run.font.name}
    bad = sorted(f for f in fonts if _SYMBOL_FONT_RE.search(f))
    checks[CheckName.NON_STANDARD_FONT] = (
        make_check(CheckName.NON_STANDARD_FONT, "fail", "Symbol/icon fonts: " + ", ".join(bad))
        if bad
        else make_check(CheckName.NON_STANDARD_FONT, "pass")
    )

    has_images = bool(doc.inline_shapes)
    text_chars = len(re.sub(r"\s", "", body_text + "".join(boxes)))
    checks[CheckName.IMAGE_BASED_PDF] = (
        make_check(CheckName.IMAGE_BASED_PDF, "fail", "Document is images only, no text")
        if has_images and text_chars < 40
        else make_check(CheckName.IMAGE_BASED_PDF, "pass")
    )

    if boxes:
        lines = boxes + lines
    return DocxExtraction(text="\n".join(lines), checks=checks)
